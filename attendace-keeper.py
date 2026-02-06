from tkinter import *
from tkinter import ttk, filedialog
from openpyxl import load_workbook
import openpyxl
import os


class GUI:
    def __init__(self, root):
        self.root = root
        self.name_font = ('Helvetica', 14, 'bold')
        self.label_font = ('Helvetica', 9, 'bold')
        self.week_var = StringVar()
        self.student_list_var = StringVar()
        self.section_var = StringVar()
        self.all_student_list = Listbox(root, width=35, height=5, selectmode='extended')
        self.attended_list = Listbox(root, width=35, height=5, selectmode='extended')
        self.section_combo = ttk.Combobox(root, textvariable=self.section_var)
        self.types_comb = ttk.Combobox(root, values=[".txt", ".xls", ".csv"], width=10)
        self.types_comb.set(".txt")
        self.week_entry = Entry(root, textvariable=self.week_var)

        self.file_importer = FileImporter()
        self.section_combo.bind("<<ComboboxSelected>>", self.on_section_selected)
        self.populate_section_combobox()
        self.add_remove_manager = AddRemoveManager(self)
        self.file_exporter = FileExporter(self, self.file_importer)

        self.create_widgets()


    def create_widgets(self):
        
        name = Label(self.root, text="AttendanceKeeper v1.0", font=self.name_font)
        select_label1 = Label(self.root, text="Select student list Excel file: ", font=self.label_font)
        self.import_button = Button(self.root, text="Import List", width=20, command=self.import_list)
        select_label2 = Label(self.root, text="Select a Student: ", font=self.label_font)
        section_label = Label(self.root, text="Section:", font=self.label_font)
        attended_label = Label(self.root, text="Attended Students:", font=self.label_font)

        name.grid(row=0, column=0, columnspan=5, pady=10)
        select_label1.grid(row=1, column=0, columnspan=2)
        self.import_button.grid(row=1, column=2, padx=10)
        select_label2.grid(row=2, column=0, columnspan=2)
        section_label.grid(row=2, column=2)
        attended_label.grid(row=2, column=3, columnspan=2)

        self.all_student_list.grid(row=3, column=0,rowspan=3, columnspan=2, padx=5)
        scrollbar_all = Scrollbar(self.root, orient=VERTICAL, command=self.all_student_list.yview)
        scrollbar_all.grid(row=3, column=1, rowspan=3, sticky="nse")
        self.section_combo.grid(row=3, column=2)
        add_button = Button(self.root, text="Add =>", width=20, command=self.add_remove_manager.add_selected_std)
        add_button.grid(row=4, column=2)
        remove_button = Button(self.root, text="<= Remove", width=20, command=self.add_remove_manager.remove_selected_std)
        remove_button.grid(row=5, column=2)
        self.attended_list.grid(row=3, column=3, rowspan=3,columnspan=2, padx=5)
        scrollbar_attended = Scrollbar(self.root, orient=VERTICAL, command=self.attended_list.yview)
        scrollbar_attended.grid(row=3, column=4, rowspan=3, sticky="nse", padx=2)

        select_type_label = Label(self.root, text="Please select file type: ")
        select_type_label.grid(row=6, column=0)
        self.types_comb.grid(row=6, column=1)
        enter_week_label = Label(self.root, text="Please enter week: ")
        enter_week_label.grid(row=6, column=2)
        self.week_entry.grid(row=6, column=3)
        self.export_button = Button(self.root, text="Export as File", command=self.export_attendance)
        self.export_button.grid(row=6, column=4, pady=5)

        self.all_student_list.config(yscrollcommand=scrollbar_all.set)
        self.attended_list.config(yscrollcommand=scrollbar_attended.set)
    
    def import_list (self):

        filename = self.file_importer.prompt_file_dialog()
        if filename:
            data = self.file_importer.load_data()
            result_dict = self.file_importer.sort_data(data)

            if result_dict:
                dict_items = list(result_dict.items())
                dict_items.pop(-1)
                result_dict = dict(dict_items)
                self.populate_section_combobox()
    
    def populate_section_combobox(self):
        section_keys = self.file_importer.get_section_keys()
        self.section_combo['values'] = section_keys
        if section_keys:
            self.section_var.set(section_keys[0])
            self.populate_student_list()

############### methods when section is changed #############
    def on_section_selected(self, event=None):
        self.populate_student_list()
        self.clear_attended_list()

    def populate_student_list(self):
        section = self.section_var.get()
        students = self.file_importer.get_student_list(section)
        self.all_student_list.delete(0, END)
        for student_data in students:
            self.all_student_list.insert(END, student_data)
    
    def clear_attended_list(self):# clear a list
        self.attended_list.delete(0, END)
        self.populate_student_list()
          

############# method for exporting ################
    def export_attendance (self):
        file_type = self.types_comb.get()
        if file_type == '.txt':
            self.file_exporter.export_to_txt()
        elif file_type == '.xls':
            self.file_exporter.export_to_xls()
        elif file_type == '.csv':
            self.file_exporter.export_to_csv()
    
    def repopulate_student_list (self):
        self.all_student_list.delete(0, END)
        self.populate_student_list()
        


class FileImporter:
    def __init__(self):
        self.filename = None
        self.sorted_data_dict = None
    
    def prompt_file_dialog (self):
        self.filename = filedialog.askopenfilename(filetypes=[("Excel Files", "*.xlsx")])
        return self.filename
    
    def load_data (self):
        if self.filename:
            workbook = load_workbook(self.filename)
            sheet = workbook.active
            all_data = []
            for row in sheet.iter_rows(values_only=True):
                all_data.append(row)
            return all_data
    
    def sort_data(self, data):
        data_dict = {}

        for row in data:
            key = row[-1]
            if key in data_dict:
                data_dict[key].append(list(row))
            else:
                data_dict[key] = [list(row)]
        sorted_data_dict = dict(sorted(data_dict.items()))
        self.sorted_data_dict = sorted_data_dict
        return sorted_data_dict
        
    def get_section_keys(self):
        if self.sorted_data_dict:
            section_keys = list(self.sorted_data_dict.keys())
            section_keys = [key for key in section_keys if key != 'Section']
            return section_keys
        else:
            return[]
        
    def get_student_list(self, section):
        if self.sorted_data_dict and section in self.sorted_data_dict:
            students = self.sorted_data_dict[section]
            modified_students = []
            for student_data in students:
                number, full_name, _, _ = student_data
                name_parts = full_name.split()
                surname = name_parts[-1]
                first_name = name_parts[0]
                modified_student = f"{surname}, {first_name}, {number}"
                modified_students.append(modified_student)
            modified_students.sort()
            return modified_students
        else:
            return[]
    


class AddRemoveManager:
    def __init__(self, gui_instance):
        self.gui = gui_instance

    def add_selected_std (self):
        selected_indeces = self.gui.all_student_list.curselection()
        if selected_indeces:
            for index in selected_indeces[::-1]:
                student_data = self.gui.all_student_list.get(index)
                self.gui.attended_list.insert(END, student_data)
                self.gui.all_student_list.delete(index)

    def remove_selected_std(self):
        selected_indices = self.gui.attended_list.curselection()
        if selected_indices:
            for index in selected_indices[::-1]:
                student_data = self.gui.attended_list.get(index)
                self.gui.all_student_list.insert(END, student_data)
                self.gui.attended_list.delete(index)



class FileExporter:
    def __init__ (self, gui_instance, file_importer_instance):
        self.gui = gui_instance
        self.file_data = file_importer_instance

    def filename_create (self):
        section_name = self.gui.section_combo.get()
        entry_value = self.gui.week_entry.get()
        file_type = self.gui.types_comb.get()
        cwd = os.getcwd()
        filename = f"{section_name}_{entry_value}{file_type}"
        file_path = os.path.join(cwd, filename)
        return file_path

    def format_std_data(self):
        file_type = self.gui.types_comb.get()
        attended_students = self.gui.attended_list.get(0, END)
        all_students = self.file_data.load_data()
        formatted_students = []
        for student in attended_students:
            surname, name, student_id = student.split(", ")
            department = ""
            for std in all_students:
                if std[0] == int(student_id.strip()):
                    department = std[2]
                    break
            if file_type == '.txt':
                formatted_student = f"{student_id}\t{name} {surname}\t{department}\n"
            elif file_type == ".xls":
                formatted_student = [student_id, f"{name} {surname}", department]
            formatted_students.append(formatted_student)
        return formatted_students
         
    def export_to_txt(self):
        try:
            filename = self.filename_create()
            data = self.format_std_data()
            header = "ID\tName\tDepartment\n"

            with open(filename, "w", encoding="utf-8") as file:
                file.write(header)
                for student in data:
                    file.write(student)
            self.gui.attended_list.delete(0, END)
            self.gui.repopulate_student_list()
            self.gui.week_var.set("")
        except Exception as e:
            print(f"Error: {e}")

    def export_to_xls(self):
        try:
            filename = self.filename_create()
            data = self.format_std_data()
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(["ID", "Name", "Department"])
            for student in data:
                sheet.append(student)
            workbook.save(filename)
            self.gui.attended_list.delete(0, END)
            self.gui.repopulate_student_list()
            self.gui.week_var.set("")
        except Exception as e:
            print(f"Error: {e}")

    def export_to_csv(self):
        raise BaseException("File type is not supported!")




if __name__ == '__main__':
    root = Tk()
    gui = GUI(root)
    gui.create_widgets()
    file_importer = FileImporter()
    add_remove_manager = AddRemoveManager(gui)
    file_exporter = FileExporter(gui, file_importer)
    root.mainloop()